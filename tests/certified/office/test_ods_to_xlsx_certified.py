"""PR-A1: Certified ODS to XLSX converter tests.

Verifies that ODS->XLSX conversion produces a real, valid XLSX file
with the correct content integrity.
"""

from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.core.settings import settings
from app.main import app


def test_ods_to_xlsx_conversion_creates_valid_xlsx(tmp_path: Path):
    client = TestClient(app)

    ods_path = Path("tests/sample.ods")
    assert ods_path.exists(), "Sample ODS is missing"

    resp = client.post(
        "/convert",
        data={"target_format": "xlsx", "operation": "ods-to-xlsx"},
        files={
            "file": (
                ods_path.name,
                ods_path.read_bytes(),
                "application/vnd.oasis.opendocument.spreadsheet",
            )
        },
    )

    assert resp.status_code == 201, resp.text
    payload = resp.json()
    assert payload.get("status") == "success"
    download_path = payload.get("download_path")
    assert download_path, payload
    assert download_path.startswith("/download/")
    relative_parts = Path(download_path.removeprefix("/download/")).parts
    assert len(relative_parts) == 2, f"Unexpected download path shape: {download_path}"
    conversion_id, filename = relative_parts
    assert filename.endswith(".xlsx")

    local_path = settings.OUTPUT_DIR / conversion_id / filename
    download_resp = client.get(download_path)
    assert download_resp.status_code == 200, download_resp.text
    assert download_resp.content, "Downloaded content is empty"

    assert local_path.exists(), f"Expected output XLSX not found: {local_path}"
    assert local_path.stat().st_size > 0, "Output XLSX is empty"
    assert local_path.suffix.lower() == ".xlsx"

    assert len(list((settings.OUTPUT_DIR / conversion_id).glob("*"))) == 1


def test_ods_to_xlsx_content_integrity(tmp_path: Path):
    """Verify the output XLSX contains the data from the source ODS."""
    client = TestClient(app)

    ods_path = Path("tests/sample.ods")
    resp = client.post(
        "/convert",
        data={"target_format": "xlsx", "operation": "ods-to-xlsx"},
        files={
            "file": (
                ods_path.name,
                ods_path.read_bytes(),
                "application/vnd.oasis.opendocument.spreadsheet",
            )
        },
    )
    assert resp.status_code == 201, resp.text
    payload = resp.json()
    conversion_id, filename = Path(payload["download_path"].removeprefix("/download/")).parts
    local_path = settings.OUTPUT_DIR / conversion_id / filename

    workbook = load_workbook(str(local_path))
    try:
        assert "Sheet1" in workbook.sheetnames, f"Expected Sheet1, got {workbook.sheetnames}"
        ws = workbook["Sheet1"]

        rows = list(ws.iter_rows(values_only=True))
        assert len(rows) >= 2, "Expected at least 2 rows (header + data)"
        header = list(rows[0])
        assert "Name" in header, f"Expected 'Name' in header, got {header}"
        assert "Value" in header, f"Expected 'Value' in header, got {header}"

        data_values = [str(cell) for row in rows[1:] for cell in row if cell is not None]
        assert "Sample ODS" in data_values, f"Expected 'Sample ODS' in data, got {data_values}"
        assert "123" in data_values, f"Expected '123' in data, got {data_values}"
    finally:
        workbook.close()
        local_path.unlink(missing_ok=True)


def test_ods_to_xlsx_output_is_real_ooxml_xlsx():
    """Verify the output starts with the OOXML ZIP magic bytes (PK)."""
    client = TestClient(app)

    ods_path = Path("tests/sample.ods")
    resp = client.post(
        "/convert",
        data={"target_format": "xlsx", "operation": "ods-to-xlsx"},
        files={
            "file": (
                ods_path.name,
                ods_path.read_bytes(),
                "application/vnd.oasis.opendocument.spreadsheet",
            )
        },
    )
    assert resp.status_code == 201, resp.text
    payload = resp.json()
    conversion_id, filename = Path(payload["download_path"].removeprefix("/download/")).parts
    local_path = settings.OUTPUT_DIR / conversion_id / filename

    try:
        with local_path.open("rb") as handle:
            header = handle.read(4)
        assert header == b"PK\x03\x04", f"Output is not a real XLSX (ZIP) file: {header!r}"
    finally:
        local_path.unlink(missing_ok=True)