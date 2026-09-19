"""Certified XLS to XLSX converter tests.

Verifies that XLS->XLSX conversion produces a real, valid OOXML XLSX file with
correct content integrity, and that non-genuine legacy files (a renamed OOXML
file and a truncated BIFF8 file) are honestly rejected rather than silently
converted.
"""

import shutil
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.core.settings import settings
from app.main import app

ASSETS = Path(__file__).resolve().parents[2] / "assets" / "regression"
SAMPLE = ASSETS / "sample_legacy_biff8.xls"
MASQUERADE = ASSETS / "masquerade_xlsx_named_xls.xls"
TRUNCATED = ASSETS / "corrupt_truncated_biff8.xls"

requires_libreoffice = pytest.mark.skipif(
    shutil.which(settings.XLS_CONVERTER_SOFFICE_PATH) is None,
    reason="LibreOffice (soffice) is not installed; XLS->XLSX requires it.",
)


def _convert(client: TestClient, sample_path: Path):
    return client.post(
        "/convert",
        data={"target_format": "xlsx", "operation": "xls-to-xlsx"},
        files={
            "file": (
                sample_path.name,
                sample_path.read_bytes(),
                "application/vnd.ms-excel",
            )
        },
    )


@requires_libreoffice
def test_xls_to_xlsx_conversion_creates_valid_xlsx(tmp_path: Path):
    client = TestClient(app)

    assert SAMPLE.exists(), "Sample legacy XLS is missing"

    resp = _convert(client, SAMPLE)

    assert resp.status_code == 201, resp.text
    payload = resp.json()
    assert payload.get("status") == "success"
    download_path = payload.get("download_path")
    assert download_path, payload
    assert download_path.startswith("/download/")
    relative_parts = Path(download_path.removeprefix("/download/")).parts
    assert len(relative_parts) == 2, f"Unexpected download path shape: {download_path}"
    conversion_id, filename = relative_parts
    assert filename.endswith(".xlsx")

    local_path = settings.OUTPUT_DIR / conversion_id / filename
    download_resp = client.get(download_path)
    assert download_resp.status_code == 200, download_resp.text
    assert download_resp.content, "Downloaded content is empty"

    assert local_path.exists(), f"Expected output XLSX not found: {local_path}"
    assert local_path.stat().st_size > 0, "Output XLSX is empty"
    assert local_path.suffix.lower() == ".xlsx"

    assert len(list((settings.OUTPUT_DIR / conversion_id).glob("*"))) == 1


@requires_libreoffice
def test_xls_to_xlsx_content_integrity(tmp_path: Path):
    """Verify the output XLSX preserves the source workbook's sheets and headers."""
    client = TestClient(app)

    resp = _convert(client, SAMPLE)
    assert resp.status_code == 201, resp.text
    payload = resp.json()
    conversion_id, filename = Path(payload["download_path"].removeprefix("/download/")).parts
    local_path = settings.OUTPUT_DIR / conversion_id / filename

    workbook = load_workbook(str(local_path))
    try:
        for expected_sheet in ("Ledger", "Summary", "Formats", "Emptyish"):
            assert expected_sheet in workbook.sheetnames, (
                f"Expected sheet {expected_sheet!r}, got {workbook.sheetnames}"
            )

        ledger = workbook["Ledger"]
        header = [str(cell) for cell in next(ledger.iter_rows(min_row=4, max_row=4, values_only=True))]
        assert header[:6] == [
            "row_id",
            "label",
            "quantity",
            "unit_price",
            "amount",
            "tax_rate",
        ], f"Unexpected Ledger header: {header}"

        assert ledger.max_row >= 100, f"Expected >=100 rows preserved, got {ledger.max_row}"
    finally:
        workbook.close()
        local_path.unlink(missing_ok=True)


@requires_libreoffice
def test_xls_to_xlsx_output_is_real_ooxml_xlsx():
    """Verify the output starts with the OOXML ZIP magic bytes (PK)."""
    client = TestClient(app)

    resp = _convert(client, SAMPLE)
    assert resp.status_code == 201, resp.text
    payload = resp.json()
    conversion_id, filename = Path(payload["download_path"].removeprefix("/download/")).parts
    local_path = settings.OUTPUT_DIR / conversion_id / filename

    try:
        with local_path.open("rb") as handle:
            header = handle.read(4)
        assert header == b"PK\x03\x04", f"Output is not a real XLSX (ZIP) file: {header!r}"
    finally:
        local_path.unlink(missing_ok=True)


def test_xls_to_xlsx_rejects_ooxml_masquerading_as_xls():
    """A genuine OOXML file renamed to .xls must be refused, never converted."""
    client = TestClient(app)

    assert MASQUERADE.exists(), "Masquerade fixture is missing"
    resp = _convert(client, MASQUERADE)

    assert resp.status_code != 201, "OOXML masquerading as XLS was accepted"
    payload = resp.json()
    assert payload.get("status") != "success"
    assert payload.get("code") == "CONVERSION_FAILED", payload
    assert payload.get("message"), "Expected error message in response"


def test_xls_to_xlsx_rejects_truncated_biff8():
    """A truncated BIFF8 file must be reported as corrupt, not converted."""
    client = TestClient(app)

    assert TRUNCATED.exists(), "Truncated fixture is missing"
    resp = _convert(client, TRUNCATED)

    assert resp.status_code != 201, "Truncated BIFF8 file was accepted"
    payload = resp.json()
    assert payload.get("status") != "success"
    assert payload.get("code") == "CONVERSION_FAILED", payload
    assert payload.get("message"), "Expected error message in response"


def test_xls_to_xlsx_plugin_supports_guard():
    """The plugin only advertises the xls->xlsx pair."""
    from app.plugins.document.xls_to_xlsx import XlsToXlsxPlugin

    plugin = XlsToXlsxPlugin()
    assert plugin.supports("xls", "xlsx") is True
    assert plugin.supports("xlsx", "xls") is False
    assert plugin.supports("xls", "pdf") is False
    assert plugin.supports("docx", "xlsx") is False
