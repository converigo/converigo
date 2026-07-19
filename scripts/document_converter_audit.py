import asyncio
import json
from pathlib import Path

from app.plugins.registry import registry as plugin_registry

ROOT = Path('tests/assets/regression')
OUT = Path('build')
OUT.mkdir(exist_ok=True)

cases = [
    {"name":"DOCX -> PDF","source":ROOT / 'sample.docx',"target":"pdf","type":"docx"},
    {"name":"XLSX -> PDF","source":ROOT / 'sample.xlsx',"target":"pdf","type":"xlsx"},
    {"name":"PDF -> DOCX","source":ROOT / 'sample.pdf',"target":"docx","type":"pdf"},
]

results = []

async def validate_input(case):
    src = case['source']
    info = {'exists': src.exists(), 'size': None, 'can_open': False, 'notes': None}
    if not src.exists():
        info['notes']='missing'
        return info
    info['size']=src.stat().st_size
    t = case['type']
    try:
        if t=='docx':
            try:
                import docx
            except ImportError:
                info['notes']='python-docx not installed'
                return info
            doc = docx.Document(str(src))
            info['can_open']=True
        elif t=='xlsx':
            try:
                import openpyxl
            except ImportError:
                info['notes']='openpyxl not installed'
                return info
            wb = openpyxl.load_workbook(str(src), read_only=True)
            info['can_open']=True
        elif t=='pdf':
            try:
                import fitz
            except ImportError:
                info['notes']='PyMuPDF not installed'
                return info
            doc = fitz.open(str(src))
            info['can_open']=True
            info['pages']=doc.page_count
    except Exception as e:
        info['notes']=str(e)
    return info

async def run_case(case):
    src = case['source']
    tgt = case['target']
    name = case['name']
    entry = {'converter':name, 'input':str(src), 'input_check':None, 'output':None, 'result':None, 'error':None}
    input_info = await validate_input(case)
    entry['input_check']=input_info
    if not input_info.get('exists') or not input_info.get('can_open'):
        entry['result']='FAIL'
        entry['error']='Input fixture missing or cannot be opened: ' + str(input_info.get('notes'))
        return entry
    try:
        plugin = plugin_registry.get_plugin(src.suffix.lstrip('.'), tgt)
    except Exception as e:
        entry['result']='FAIL'
        entry['error']='Plugin resolution failed: '+str(e)
        return entry
    try:
        out_path = await plugin.convert(src, tgt)
        entry['output']=str(out_path)
        if not out_path.exists() or out_path.stat().st_size==0:
            entry['result']='FAIL'
            entry['error']='Output missing or empty'
            return entry
        # validate output
        if tgt=='pdf':
            try:
                import fitz
                doc = fitz.open(str(out_path))
                if doc.page_count>0:
                    entry['result']='PASS'
                else:
                    entry['result']='FAIL'
                    entry['error']='PDF has zero pages'
            except ImportError:
                # fallback to size check
                if out_path.stat().st_size>0:
                    entry['result']='PASS'
                else:
                    entry['result']='FAIL'
                    entry['error']='PyMuPDF not available and file empty'
            except Exception as e:
                entry['result']='FAIL'
                entry['error']='PyMuPDF open failed: '+str(e)
        elif tgt=='docx':
            try:
                import docx
                doc = docx.Document(str(out_path))
                entry['result']='PASS'
            except ImportError:
                entry['result']='FAIL'
                entry['error']='python-docx not installed'
            except Exception as e:
                entry['result']='FAIL'
                entry['error']='python-docx open failed: '+str(e)
        elif tgt=='xlsx':
            try:
                import openpyxl
                wb = openpyxl.load_workbook(str(out_path), read_only=True)
                entry['result']='PASS'
            except ImportError:
                entry['result']='FAIL'
                entry['error']='openpyxl not installed'
            except Exception as e:
                entry['result']='FAIL'
                entry['error']='openpyxl open failed: '+str(e)
        else:
            entry['result']='PASS'
    except Exception as e:
        entry['result']='FAIL'
        entry['error']=str(e)
    return entry

async def main():
    for c in cases:
        print('Running', c['name'])
        r = await run_case(c)
        results.append(r)
        print('Result', r['result'])
    with open(OUT/'document_converter_audit_results.json','w',encoding='utf8') as fh:
        json.dump(results, fh, indent=2, ensure_ascii=False)

    # write human report
    lines = ['# DOCUMENT CONVERTER AUDIT REPORT', '']
    for r in results:
        lines.append('Converter: '+r['converter'])
        lines.append('Input: '+r['input'])
        lines.append('Input check: '+str(r['input_check']))
        lines.append('Engine: see plugin metadata')
        lines.append('Result: '+str(r['result']))
        lines.append('Error: '+str(r['error']))
        lines.append('')

    with open('DOCUMENT_CONVERTER_AUDIT_REPORT.md','w',encoding='utf8') as fh:
        fh.write('\n'.join(lines))

    print('Wrote build/document_converter_audit_results.json and DOCUMENT_CONVERTER_AUDIT_REPORT.md')

if __name__=='__main__':
    asyncio.run(main())
